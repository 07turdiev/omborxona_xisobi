"""Ma'lumotni Excel faylga chiqarish.

Sonlar matn emas, **son** sifatida yoziladi: aks holda Excel ularni
qo'sha olmaydi va har safar qo'lda formatlash kerak bo'lardi.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='F1EDFF')
HEADER_FONT = Font(bold=True, size=10)

# Ustun turlari — format va kenglikni belgilaydi
MONEY = 'money'
NUMBER = 'number'
DATE = 'date'
DATETIME = 'datetime'
TEXT = 'text'

FORMATS = {
    MONEY: '#,##0.00',
    NUMBER: '#,##0',
    DATE: 'DD.MM.YYYY',
    DATETIME: 'DD.MM.YYYY HH:MM',
}

WIDTHS = {MONEY: 16, NUMBER: 10, DATE: 12, DATETIME: 18, TEXT: 22}

NUMERIC_KINDS = frozenset({MONEY, NUMBER})


class Column:
    """Chiqariladigan ustun: `key` — qatordagi kalit yoki atribut."""

    __slots__ = ('key', 'title', 'kind', 'width')

    def __init__(self, key: str, title: str, kind: str = TEXT, width: int | None = None):
        if kind not in WIDTHS:
            raise ValueError(f'Noma\'lum ustun turi: {kind!r}')

        self.key = key
        self.title = title
        self.kind = kind
        self.width = width or WIDTHS[kind]


def _value_of(row: Any, key: str) -> Any:
    """Qatordan qiymat: lug'at ham, obyekt ham bo'lishi mumkin ("a.b.c")."""
    value = row

    for part in key.split('.'):
        if value is None:
            return None

        value = value.get(part) if isinstance(value, dict) else getattr(value, part, None)

    return value


def _cell_value(value: Any, kind: str) -> Any:
    """Qiymatni Excel tushunadigan tipga o'giradi."""
    if value is None:
        return None

    if isinstance(value, bool):
        return 'ha' if value else 'yo‘q'

    if isinstance(value, str) and kind in NUMERIC_KINDS:
        # DRF `Decimal` ni matn sifatida beradi ("1500.00")
        try:
            return float(Decimal(value))
        except (InvalidOperation, ValueError):
            return value

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, datetime):
        # Excel vaqt zonasini bilmaydi — mahalliy vaqtga o'giramiz
        if timezone.is_aware(value):
            value = timezone.localtime(value)

        return value.replace(tzinfo=None)

    if isinstance(value, date):
        return value

    return value


def add_sheet(
    workbook: Workbook,
    columns: Sequence[Column],
    rows: Iterable[Any],
    *,
    sheet_name: str = 'Hisobot',
    title: str = '',
    sheet: Any = None,
) -> Any:
    """Kitobga bitta varaq qo'shadi."""
    if sheet is None:
        sheet = workbook.create_sheet()

    # Excel varaq nomida ba'zi belgilarni qabul qilmaydi
    sheet.title = sheet_name[:31].replace('/', '-').replace('\\', '-')

    offset = 0

    if title:
        sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        offset = 2

    header_row = offset + 1

    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=column.title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)

        sheet.column_dimensions[get_column_letter(index)].width = column.width

    for row_index, row in enumerate(rows, start=header_row + 1):
        for column_index, column in enumerate(columns, start=1):
            cell = sheet.cell(
                row=row_index,
                column=column_index,
                value=_cell_value(_value_of(row, column.key), column.kind),
            )

            if column.kind in FORMATS:
                cell.number_format = FORMATS[column.kind]

    # Sarlavha doim ko'rinib tursin va filtr ishlasin
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(len(columns))}{sheet.max_row}'

    return sheet


def build_workbook(
    columns: Sequence[Column],
    rows: Iterable[Any],
    *,
    sheet_name: str = 'Hisobot',
    title: str = '',
) -> Workbook:
    """Bitta jadvalni Excel kitobiga aylantiradi."""
    workbook = Workbook()

    add_sheet(
        workbook, columns, rows, sheet_name=sheet_name, title=title, sheet=workbook.active
    )

    return workbook


def excel_response(workbook: Workbook, filename: str) -> HttpResponse:
    """Kitobni yuklab olinadigan javobga o'raydi."""
    stamp = timezone.localdate().strftime('%Y-%m-%d')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    # Fayl nomida o'zbekcha harflar bo'lishi mumkin — RFC 5987
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}-{stamp}.xlsx"

    workbook.save(response)

    return response
