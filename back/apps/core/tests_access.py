"""Bo'lim ruxsatlari va moliyaviy maydonlarni yashirish testlari.

API testlari haqiqiy JWT bilan ishlaydi: `TenantMiddleware` tokenni DRF'dan
oldin o'qib, tenantni shundan oladi, `force_authenticate` esa buni chetlab
o'tadi va test noto'g'ri sababdan o'tib ketardi.
"""

from __future__ import annotations

import io
from decimal import Decimal

from django.test import SimpleTestCase, TestCase, tag
from openpyxl import load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from apps.catalog.models import Category, Product, Variant
from apps.core.access import (
    ALL_PERMISSIONS,
    ROLE_DEFAULTS,
    Perm,
    hidden_keys,
    redact,
)
from apps.core.tenancy import tenant_context
from apps.documents.models import Document
from apps.tenants.models import Membership, Tenant
from apps.users.models import User
from apps.warehouse.models import Warehouse

Role = Membership.Role


# ---------------------------------------------------------------------
# Model: amaldagi ruxsatlar
# ---------------------------------------------------------------------

@tag('access')
class EffectivePermissionTests(SimpleTestCase):

    def test_owner_always_has_everything(self):
        """Egasini hech kim ruxsatsiz qoldira olmaydi — hatto bo'sh ro'yxat bilan."""
        membership = Membership(role=Role.OWNER, is_active=True, permissions=[])

        self.assertEqual(membership.effective_permissions, ALL_PERMISSIONS)

    def test_role_defaults_when_not_customized(self):
        membership = Membership(role=Role.STOREKEEPER, is_active=True)

        self.assertEqual(membership.effective_permissions, ROLE_DEFAULTS['storekeeper'])
        self.assertTrue(membership.uses_role_defaults)

    def test_custom_list_replaces_role_defaults(self):
        membership = Membership(
            role=Role.STOREKEEPER, is_active=True, permissions=[Perm.SALES]
        )

        self.assertTrue(membership.has_perm(Perm.SALES))
        self.assertFalse(membership.has_perm(Perm.IMPORTS))
        self.assertFalse(membership.uses_role_defaults)

    def test_unknown_codes_ignored(self):
        membership = Membership(
            role=Role.VIEWER, is_active=True, permissions=['sales', 'nonexistent']
        )

        self.assertEqual(membership.effective_permissions, {'sales'})

    def test_inactive_membership_has_nothing(self):
        membership = Membership(role=Role.OWNER, is_active=False)

        self.assertEqual(membership.effective_permissions, frozenset())

    def test_cashier_does_not_see_money_by_default(self):
        """Kassir sotadi, lekin tovar qanchaga olinganini va foydani bilmaydi."""
        membership = Membership(role=Role.SALESPERSON, is_active=True)

        self.assertTrue(membership.has_perm(Perm.SALES))
        self.assertFalse(membership.has_perm(Perm.VIEW_PURCHASE_PRICE))
        self.assertFalse(membership.has_perm(Perm.VIEW_PROFIT))

    def test_viewer_cannot_write_accountant_can(self):
        self.assertFalse(Membership(role=Role.VIEWER, is_active=True).can_write)
        self.assertTrue(Membership(role=Role.ACCOUNTANT, is_active=True).can_write)

    def test_manager_keeps_admin_sections(self):
        """Menejer avvalgidek sozlamalar va xodimlarni boshqaradi."""
        membership = Membership(role=Role.MANAGER, is_active=True)

        self.assertTrue(membership.has_all({Perm.SETTINGS, Perm.USERS}))


@tag('access')
class RedactTests(SimpleTestCase):

    def test_nested_values_become_none(self):
        data = {
            'summary': {'revenue': 100, 'gross_profit': 30},
            'rows': [{'name': 'Sement', 'cost': 70, 'profit': 30}],
        }

        result = redact(data, frozenset({'gross_profit', 'cost', 'profit'}))

        self.assertEqual(result['summary'], {'revenue': 100, 'gross_profit': None})
        self.assertEqual(result['rows'], [{'name': 'Sement', 'cost': None, 'profit': None}])

    def test_keys_are_kept_not_removed(self):
        """Kalit o'chirilsa, interfeys `undefined` ni 0 deb "foyda: 0" ko'rsatardi."""
        result = redact({'profit': 5}, frozenset({'profit'}))

        self.assertIn('profit', result)

    def test_full_access_hides_nothing(self):
        membership = Membership(role=Role.OWNER, is_active=True)

        self.assertEqual(hidden_keys(membership), frozenset())


# ---------------------------------------------------------------------
# API
# ---------------------------------------------------------------------

class AccessApiTestBase(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.tenant = Tenant.objects.create(name='Do\'kon', slug='access-api')

        with tenant_context(cls.tenant.id):
            cls.shop = Warehouse.objects.create(
                tenant=cls.tenant, code='SHOP', name='Do\'kon',
                purpose=Warehouse.Purpose.RETAIL,
            )
            category = Category.objects.create(tenant=cls.tenant, name='Sement')
            product = Product.objects.create(
                tenant=cls.tenant, category=category, name='Sement', base_unit='kg'
            )
            cls.variant = Variant.objects.create(
                tenant=cls.tenant, product=product, sku='CEM-1',
                purchase_price=Decimal('1000'), sale_price=Decimal('1300'),
            )
            cls.purchase = Document.objects.create(
                tenant=cls.tenant, kind=Document.Kind.PURCHASE, number='KIR-1',
                warehouse=cls.shop, status=Document.Status.CONFIRMED,
                total_amount=Decimal('50000'),
            )
            cls.sale = Document.objects.create(
                tenant=cls.tenant, kind=Document.Kind.SALE, number='SOT-1',
                warehouse=cls.shop, status=Document.Status.CONFIRMED,
                total_amount=Decimal('13000'), total_cost=Decimal('10000'),
            )

    _counter = 0

    def member(self, role, permissions=None):
        AccessApiTestBase._counter += 1
        user = User.objects.create_user(
            username=f'u{AccessApiTestBase._counter}-{role}', password='x'
        )

        return Membership.objects.create(
            tenant=self.tenant, user=user, role=role, permissions=permissions
        )

    def client_for(self, role, permissions=None):
        membership = self.member(role, permissions)

        client = APIClient()
        token = RefreshToken.for_user(membership.user).access_token
        client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        client.membership = membership

        return client


@tag('access')
class DocumentAccessTests(AccessApiTestBase):

    def test_cashier_cannot_see_purchase_documents(self):
        client = self.client_for(Role.SALESPERSON)

        listing = client.get('/api/documents/', {'kind': 'purchase'})
        detail = client.get(f'/api/documents/{self.purchase.pk}/')

        self.assertEqual(listing.status_code, 200)
        self.assertEqual(listing.json()['results'], [])
        self.assertEqual(detail.status_code, 404)

    def test_cashier_sees_sale_without_cost_and_profit(self):
        data = self.client_for(Role.SALESPERSON).get(f'/api/documents/{self.sale.pk}/').json()

        self.assertEqual(Decimal(data['total_amount']), Decimal('13000'))
        self.assertIsNone(data['total_cost'])
        self.assertIsNone(data['profit'])

    def test_accountant_sees_cost_and_profit(self):
        data = self.client_for(Role.ACCOUNTANT).get(f'/api/documents/{self.sale.pk}/').json()

        self.assertEqual(Decimal(data['profit']), Decimal('3000'))

    def test_purchase_amount_hidden_without_purchase_price_permission(self):
        """Kirim ruxsati bor, lekin narxni ko'rish yo'q — summa tozalanadi."""
        client = self.client_for(Role.STOREKEEPER, [Perm.IMPORTS])

        data = client.get(f'/api/documents/{self.purchase.pk}/').json()

        self.assertIsNone(data['total_amount'])

    def test_storekeeper_cannot_create_sale(self):
        response = self.client_for(Role.STOREKEEPER).post(
            '/api/documents/',
            {'kind': 'sale', 'warehouse': self.shop.pk},
            format='json',
        )

        self.assertEqual(response.status_code, 403)

    def test_custom_permission_opens_section(self):
        client = self.client_for(Role.STOREKEEPER, [Perm.SALES])

        self.assertEqual(client.get(f'/api/documents/{self.sale.pk}/').status_code, 200)
        self.assertEqual(client.get(f'/api/documents/{self.purchase.pk}/').status_code, 404)

    def test_viewer_cannot_write_even_with_permission(self):
        """Ruxsat kuzatuvchiga faqat nimani ko'rishini belgilaydi."""
        client = self.client_for(Role.VIEWER, [Perm.COUNTERPARTIES])

        response = client.post(
            '/api/partners/', {'name': 'Mijoz', 'is_customer': True}, format='json'
        )

        self.assertEqual(response.status_code, 403)


@tag('access')
class ReportAndExportAccessTests(AccessApiTestBase):

    def test_reports_require_permission(self):
        client = self.client_for(Role.SALESPERSON)

        self.assertEqual(client.get('/api/reports/summary/').status_code, 403)

    def test_dashboard_open_but_profit_hidden(self):
        response = self.client_for(Role.SALESPERSON).get('/api/reports/dashboard/')

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIsNone(data['summary']['gross_profit'])
        self.assertIsNone(data['valuation']['cost_value'])
        self.assertIsNotNone(data['summary']['revenue'])

    def test_loss_amounts_hidden_without_purchase_price(self):
        response = self.client_for(Role.VIEWER).get('/api/reports/losses/')

        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.json()['total'])

    def test_export_requires_print_permission(self):
        client = self.client_for(Role.STOREKEEPER, [Perm.STOCK])

        self.assertEqual(client.get('/api/stock/export/').status_code, 403)

    def test_export_drops_financial_columns(self):
        """Excel faylida ruxsatsiz ustun bo'sh qolmaydi — umuman chiqmaydi."""
        response = self.client_for(Role.SALESPERSON).get('/api/stock/export/')

        self.assertEqual(response.status_code, 200)

        sheet = load_workbook(io.BytesIO(response.content)).active
        cells = {cell for row in sheet.iter_rows(values_only=True) for cell in row}

        self.assertIn('Sotuv narxi', cells)
        self.assertNotIn('Tannarx (FIFO)', cells)
        self.assertNotIn('Birlik tannarxi', cells)


@tag('access')
class MembershipAccessTests(AccessApiTestBase):

    def patch(self, client, membership, payload):
        return client.patch(f'/api/members/{membership.pk}/', payload, format='json')

    def test_storekeeper_cannot_grant_warehouse_access(self):
        """Avval yozish huquqi bo'lgan istalgan xodim o'ziga ombor ochib olardi."""
        client = self.client_for(Role.STOREKEEPER)

        response = client.post(
            '/api/warehouse-access/',
            {'warehouse': self.shop.pk, 'user': client.membership.user_id, 'level': 'manage'},
            format='json',
        )

        self.assertEqual(response.status_code, 403)

    def test_manager_can_grant_warehouse_access(self):
        target = self.member(Role.STOREKEEPER)

        response = self.client_for(Role.MANAGER).post(
            '/api/warehouse-access/',
            {'warehouse': self.shop.pk, 'user': target.user_id, 'level': 'view'},
            format='json',
        )

        self.assertEqual(response.status_code, 201, response.content)

    def test_cannot_change_own_access(self):
        client = self.client_for(Role.MANAGER)

        response = self.patch(client, client.membership, {'role': Role.OWNER})

        self.assertEqual(response.status_code, 400)

    def test_non_owner_cannot_touch_owner(self):
        owner = self.member(Role.OWNER)

        response = self.patch(self.client_for(Role.MANAGER), owner, {'is_active': False})

        self.assertEqual(response.status_code, 400)

    def test_non_owner_cannot_grant_owner_role(self):
        target = self.member(Role.VIEWER)

        response = self.patch(self.client_for(Role.MANAGER), target, {'role': Role.OWNER})

        self.assertEqual(response.status_code, 400)

    def test_cannot_grant_permission_you_do_not_have(self):
        """"Xodimlar" ruxsati bor, lekin foydani ko'rmaydigan kishi uni boshqaga bera olmaydi."""
        client = self.client_for(Role.STOREKEEPER, [Perm.USERS, Perm.STOCK])
        target = self.member(Role.VIEWER)

        response = self.patch(
            client, target, {'permissions': [Perm.STOCK, Perm.VIEW_PROFIT]}
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn('view_profit', str(response.json()))

    def test_role_with_more_rights_cannot_be_assigned_by_weaker_admin(self):
        client = self.client_for(Role.STOREKEEPER, [Perm.USERS, Perm.STOCK])
        target = self.member(Role.VIEWER)

        response = self.patch(client, target, {'role': Role.MANAGER})

        self.assertEqual(response.status_code, 400)

    def test_list_equal_to_role_defaults_stored_as_null(self):
        """Keyin rol o'zgarsa, ruxsatlar ham u bilan birga o'zgarishi uchun."""
        target = self.member(Role.STOREKEEPER, [Perm.STOCK])

        response = self.patch(
            self.client_for(Role.OWNER), target,
            {'permissions': sorted(ROLE_DEFAULTS['storekeeper'])},
        )

        self.assertEqual(response.status_code, 200, response.content)
        target.refresh_from_db()
        self.assertIsNone(target.permissions)

    def test_owner_sets_custom_permissions(self):
        target = self.member(Role.SALESPERSON)

        response = self.patch(
            self.client_for(Role.OWNER), target,
            {'permissions': [Perm.SALES, Perm.VIEW_PROFIT]},
        )

        self.assertEqual(response.status_code, 200, response.content)
        target.refresh_from_db()
        self.assertEqual(target.permissions, ['sales', 'view_profit'])

    def test_unknown_permission_rejected(self):
        target = self.member(Role.VIEWER)

        response = self.patch(self.client_for(Role.OWNER), target, {'permissions': ['nope']})

        self.assertEqual(response.status_code, 400)

    def test_cannot_remove_self_or_owner(self):
        client = self.client_for(Role.MANAGER)
        owner = self.member(Role.OWNER)

        self.assertEqual(client.delete(f'/api/members/{client.membership.pk}/').status_code, 403)
        self.assertEqual(client.delete(f'/api/members/{owner.pk}/').status_code, 403)

    def test_me_includes_permissions(self):
        data = self.client_for(Role.SALESPERSON).get('/api/auth/me/').json()

        self.assertEqual(
            data['current_tenant']['permissions'], sorted(ROLE_DEFAULTS['salesperson'])
        )

    def test_roles_and_catalog_endpoints(self):
        client = self.client_for(Role.VIEWER)

        roles = client.get('/api/members/roles/').json()
        catalog = client.get('/api/members/permissions/').json()

        accountant = next(role for role in roles if role['value'] == 'accountant')
        self.assertIn('view_profit', accountant['default_permissions'])
        self.assertIn('view_profit', {item['value'] for item in catalog['items']})
