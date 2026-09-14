"""Excel import testlari.

Fayllar test ichida `openpyxl` bilan quriladi — haqiqiy foydalanuvchi
fayli qanday bo'lsa, shunday: sonlar matn sifatida, ruscha sarlavhalar,
bo'sh qatorlar.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models import Sum
from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from apps.audit.models import AuditEvent
from apps.catalog.models import Barcode, Category, Product, Variant
from apps.core.tenancy import tenant_context
from apps.dataimport.parsing import normalize_header, parse_date, parse_decimal
from apps.documents.models import Document
from apps.partners.models import Partner
from apps.stock.models import StockBalance
from apps.tenants.models import Membership, Tenant
from apps.users.models import User
from apps.warehouse.models import Warehouse

Role = Membership.Role
D = Decimal


def xlsx(headers, *rows) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))

    for row in rows:
        sheet.append(list(row))

    buffer = BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()


class ParsingTests(SimpleTestCase):

    def test_numbers(self):
        self.assertEqual(parse_decimal('7 500 000,50'), D('7500000.50'))
        self.assertEqual(parse_decimal('7,500,000.50'), D('7500000.50'))
        self.assertEqual(parse_decimal('7.500.000'), D('7500000'))
        self.assertEqual(parse_decimal("58 000 so'm"), D('58000'))
        self.assertEqual(parse_decimal(0.1), D('0.1'))
        self.assertIsNone(parse_decimal('  '))

        with self.assertRaises(ValueError):
            parse_decimal('o‘n'.replace('‘', ''))

    def test_dates(self):
        self.assertEqual(parse_date('14.09.2026'), date(2026, 9, 14))
        self.assertEqual(parse_date('2026-09-14 00:00:00'), date(2026, 9, 14))
        # Excel seriya raqami
        self.assertEqual(parse_date(46279), date(2026, 9, 14))

        with self.assertRaises(ValueError):
            parse_date('ertaga')

    def test_headers(self):
        self.assertEqual(normalize_header('Ед. изм.'), normalize_header('ед изм'))
        self.assertEqual(normalize_header('To‘lov usuli'), normalize_header("To'lov usuli"))
        self.assertEqual(normalize_header('Mahsulot nomi *'), 'mahsulot nomi')


class ImportTestBase(TestCase):

    @classmethod
    def setUpTestData(cls):
        cls.tenant = Tenant.objects.create(name='Do‘kon', slug='import-test')

        with tenant_context(cls.tenant.id):
            cls.shop = Warehouse.objects.create(
                tenant=cls.tenant, code='SHOP', name='Asosiy do‘kon',
                purpose=Warehouse.Purpose.RETAIL,
            )
            cls.root = Category.objects.create(tenant=cls.tenant, name='Qurilish mollari')
            cls.cement = Category.objects.create(
                tenant=cls.tenant, name='Sement', parent=cls.root, code_prefix='CEM',
            )
            product = Product.objects.create(
                tenant=cls.tenant, category=cls.cement, name='Sement M400', base_unit='qop',
            )
            cls.variant = Variant.objects.create(
                tenant=cls.tenant, product=product, sku='CEM-1',
                purchase_price=D('50000'), sale_price=D('58000'),
            )

    _counter = 0

    def client_for(self, role, permissions=None):
        ImportTestBase._counter += 1
        user = User.objects.create_user(username=f'imp{ImportTestBase._counter}', password='x')
        Membership.objects.create(
            tenant=self.tenant, user=user, role=role, permissions=permissions
        )

        client = APIClient()
        client.credentials(
            HTTP_AUTHORIZATION=f'Bearer {RefreshToken.for_user(user).access_token}'
        )

        return client

    def setUp(self):
        self.owner = self.client_for(Role.OWNER)

    def upload(self, kind, step, content, client=None, **extra):
        client = client or self.owner
        file = SimpleUploadedFile(
            'fayl.xlsx', content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        return client.post(
            f'/api/import/{kind}/{step}/', {'file': file, **extra}, format='multipart'
        )

    @property
    def today(self) -> str:
        return timezone.localdate().strftime('%d.%m.%Y')


class TemplateAndAccessTests(ImportTestBase):

    def test_template_is_xlsx_with_guide(self):
        response = self.owner.get('/api/import/products/template/')

        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(BytesIO(response.content))
        headers = [cell.value for cell in workbook.worksheets[0][1]]

        self.assertIn('Mahsulot nomi *', headers)
        self.assertIn('Kirim narxi', headers)
        self.assertEqual(workbook.worksheets[1].title, 'Yo‘riqnoma')

    def test_template_hides_purchase_price_without_permission(self):
        client = self.client_for(Role.STOREKEEPER, permissions=['products', 'categories'])

        response = client.get('/api/import/products/template/')
        headers = [cell.value for cell in load_workbook(BytesIO(response.content)).worksheets[0][1]]

        self.assertNotIn('Kirim narxi', headers)

    def test_section_permission_required(self):
        salesperson = self.client_for(Role.SALESPERSON)

        self.assertEqual(salesperson.get('/api/import/products/template/').status_code, 403)
        self.assertEqual(
            self.upload('purchases', 'preview', xlsx(['Ombor']), client=salesperson).status_code,
            403,
        )

    def test_viewer_cannot_import(self):
        viewer = self.client_for(Role.VIEWER)

        response = self.upload('categories', 'preview', xlsx(['Kategoriya nomi'], ['X']), client=viewer)

        self.assertEqual(response.status_code, 403)

    def test_unknown_type_and_bad_file(self):
        self.assertEqual(self.owner.get('/api/import/nimadir/template/').status_code, 404)

        response = self.upload('categories', 'preview', b'bu excel emas')

        self.assertEqual(response.status_code, 400)
        self.assertIn('xlsx', response.json()['detail'])


class CategoryImportTests(ImportTestBase):

    def test_creates_tree_and_is_idempotent(self):
        content = xlsx(
            ['Kategoriya nomi', 'Yuqori kategoriya', 'Ostki kategoriyalar', 'Kod prefiksi'],
            ['Bo‘yoqlar', 'Qurilish mollari > Pardozlash', 'Emal; Suv emulsiya', 'BY'],
            ['sement', 'Qurilish mollari', '', 'CMT'],
        )

        preview = self.upload('categories', 'preview', content).json()

        self.assertTrue(preview['can_commit'], preview)
        self.assertEqual(preview['summary']['create'], 1)
        # Mavjud "Sement" — katta-kichik harfdan qat'i nazar topiladi, prefiksi yangilanadi
        self.assertEqual(preview['summary']['update'], 1)
        self.assertIn('Pardozlash', preview['rows'][0]['warnings'][0])

        response = self.upload('categories', 'commit', content)
        self.assertEqual(response.status_code, 201, response.content)

        with tenant_context(self.tenant.id):
            paint = Category.objects.get(name='Bo‘yoqlar')
            self.assertEqual(paint.parent.name, 'Pardozlash')
            self.assertEqual(paint.parent.parent_id, self.root.pk)
            self.assertEqual(paint.children.count(), 2)
            self.cement.refresh_from_db()
            self.assertEqual(self.cement.code_prefix, 'CMT')
            self.assertEqual(
                AuditEvent.objects.filter(action='import', object_type='category').count(), 1
            )

        # Shu fayl qayta: avval import qilingani aytiladi
        again = self.upload('categories', 'preview', content).json()
        self.assertIsNotNone(again['duplicate'])

        self.assertEqual(self.upload('categories', 'commit', content).status_code, 409)

        # Ruxsat berilsa ham o'zgarish yo'q — hammasi mavjud
        forced = self.upload('categories', 'commit', content, allow_duplicate='true')
        self.assertEqual(forced.status_code, 400)


class ProductImportTests(ImportTestBase):

    def test_russian_headers_create_and_update(self):
        content = xlsx(
            ['Компания', 'Категория', 'Подкатегория', 'Артикул', 'Название товара',
             'Штрих-код', 'Ед. изм.', 'Закупочная цена', 'Цена продажи', 'Статус'],
            ['UNISER', 'Qurilish mollari', 'Sement', '', 'Sement M500', 4780000000011,
             'qop', '61 000', '69000', 'Активен'],
            ['UNISER', 'Qurilish mollari > Sement', '', 'CEM-1', 'Sement M400',
             '', '', '', '60 000,00', ''],
        )

        preview = self.upload('products', 'preview', content).json()

        self.assertEqual(preview['invalid'], 0, preview['rows'])
        self.assertEqual(preview['summary'], {'create': 1, 'update': 1, 'skip': 0, 'documents': 0})
        self.assertEqual(preview['unknown_columns'], [])

        response = self.upload('products', 'commit', content)
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['created'], 1)

        with tenant_context(self.tenant.id):
            variant = Variant.objects.select_related('product').get(product__name='Sement M500')
            # Artikul kategoriya prefiksidan
            self.assertTrue(variant.sku.startswith('CEM-'))
            self.assertEqual(variant.purchase_price, D('61000'))
            self.assertEqual(
                Barcode.objects.get(variant=variant).code_type, Barcode.CodeType.EAN13
            )

            self.variant.refresh_from_db()
            self.assertEqual(self.variant.sale_price, D('60000'))
            # Bo'sh katak mavjud qiymatni o'chirmaydi
            self.assertEqual(self.variant.purchase_price, D('50000'))

    def test_errors_block_whole_file(self):
        content = xlsx(
            ['Kategoriya', 'Mahsulot nomi', 'Artikul', 'Sotuv narxi'],
            ['Qurilish mollari', 'Yaxshi mahsulot', 'OK-1', '100'],
            ['Yo‘q kategoriya', 'Yomon mahsulot', 'BAD-1', 'yuz'],
            ['Qurilish mollari', 'Takror', 'OK-1', '100'],
        )

        preview = self.upload('products', 'preview', content).json()

        self.assertFalse(preview['can_commit'])
        self.assertEqual(preview['invalid'], 2)
        errors = ' '.join(preview['rows'][1]['errors'])
        self.assertIn('Kategoriya topilmadi', errors)
        self.assertIn('son emas', errors)
        self.assertIn('2-qatorda ham bor', preview['rows'][2]['errors'][0])

        self.assertEqual(self.upload('products', 'commit', content).status_code, 400)

        with tenant_context(self.tenant.id):
            self.assertFalse(Variant.objects.filter(sku='OK-1').exists())

    def test_missing_required_column(self):
        preview = self.upload('products', 'preview', xlsx(['Artikul'], ['X'])).json()

        self.assertIn('Mahsulot nomi', preview['missing_columns'])
        self.assertFalse(preview['can_commit'])

    def test_purchase_price_column_ignored_without_permission(self):
        client = self.client_for(Role.STOREKEEPER, permissions=['products', 'categories'])
        content = xlsx(
            ['Kategoriya', 'Mahsulot nomi', 'Artikul', 'Kirim narxi'],
            ['Qurilish mollari > Sement', 'Maxfiy', 'SEC-1', '1000'],
        )

        preview = self.upload('products', 'preview', content, client=client).json()

        self.assertEqual(preview['hidden_columns'], ['Kirim narxi'])
        self.assertNotIn('purchase_price', preview['rows'][0]['values'])

        self.assertEqual(self.upload('products', 'commit', content, client=client).status_code, 201)

        with tenant_context(self.tenant.id):
            self.assertIsNone(Variant.objects.get(sku='SEC-1').purchase_price)

    def test_base_unit_locked_after_movements(self):
        self.upload('purchases', 'commit', xlsx(
            ['Ombor', 'Sana', 'Artikul', 'Miqdor', 'Kirim narxi'],
            ['SHOP', self.today, 'CEM-1', 5, 50000],
        ))

        preview = self.upload('products', 'preview', xlsx(
            ['Kategoriya', 'Mahsulot nomi', 'Artikul', 'Birlik'],
            ['Qurilish mollari > Sement', 'Sement M400', 'CEM-1', 'kg'],
        )).json()

        self.assertIn('birligini o‘zgartirib bo‘lmaydi', preview['rows'][0]['errors'][0])


class PurchaseImportTests(ImportTestBase):

    def test_groups_rows_into_confirmed_documents(self):
        content = xlsx(
            ['Ombor', 'Sana', 'Yetkazib beruvchi', 'Yetkazib beruvchi hujjati', 'To‘lov usuli',
             'Artikul', 'Miqdor', 'Kirim narxi', 'Sotuv narxi', 'Partiya', 'Yaroqlilik muddati'],
            ['SHOP', self.today, 'Yangi Ta’minot MChJ', 'NK-1', 'Перечисление',
             'CEM-1', '100', '50 000', '59000', 'P-01', '01.01.2027'],
            ['Asosiy do‘kon', self.today, 'yangi ta’minot mchj', 'NK-1', 'o‘tkazma',
             'CEM-1', 20, 50000, '', 'P-01', '01.01.2027'],
            ['SHOP', self.today, '', '', '', 'CEM-1', 5, 51000, '', '', ''],
        )

        preview = self.upload('purchases', 'preview', content).json()

        self.assertTrue(preview['can_commit'], preview['rows'])
        self.assertEqual(preview['summary']['documents'], 2)
        self.assertIn('Yangi yetkazib beruvchi', preview['rows'][0]['warnings'][0])

        response = self.upload('purchases', 'commit', content)
        self.assertEqual(response.status_code, 201, response.content)

        body = response.json()
        self.assertEqual(body['created'], 2)
        self.assertEqual({doc['status'] for doc in body['documents']}, {'confirmed'})

        with tenant_context(self.tenant.id):
            total = StockBalance.objects.filter(variant=self.variant, warehouse=self.shop).aggregate(
                total=Sum('quantity')
            )['total']
            self.assertEqual(total, D('125'))

            first = Document.objects.get(pk=body['documents'][0]['id'])
            self.assertEqual(first.lines.count(), 2)
            self.assertEqual(first.external_number, 'NK-1')
            self.assertEqual(first.payment_method, Document.PaymentMethod.TRANSFER)
            self.assertEqual(Partner.objects.filter(name='Yangi Ta’minot MChJ').count(), 1)

            self.variant.refresh_from_db()
            self.assertEqual(self.variant.sale_price, D('59000'))

            self.assertEqual(AuditEvent.objects.filter(action='import', object_type='purchase').count(), 2)

        # Shu nakladnoy boshqa faylda qayta kelsa — ogohlantirish
        repeat = self.upload('purchases', 'preview', xlsx(
            ['Ombor', 'Sana', 'Yetkazib beruvchi', 'Yetkazib beruvchi hujjati', 'Artikul',
             'Miqdor', 'Kirim narxi'],
            ['SHOP', self.today, 'Yangi Ta’minot MChJ', 'NK-1', 'CEM-1', 1, 50000],
        )).json()
        self.assertIn('avval kiritilgan', ' '.join(repeat['rows'][0]['warnings']))

    def test_draft_mode_and_row_checks(self):
        content = xlsx(
            ['Ombor', 'Sana', 'Artikul', 'Miqdor', 'Kirim narxi'],
            ['SHOP', self.today, 'CEM-1', 3, 50000],
        )

        response = self.upload('purchases', 'commit', content, confirm='false')
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['documents'][0]['status'], 'draft')

        bad = self.upload('purchases', 'preview', xlsx(
            ['Ombor', 'Sana', 'Artikul', 'Miqdor', 'Kirim narxi', 'To‘lov usuli'],
            ['OMBOR-YOQ', '01.01.2099', 'YOQ-SKU', 0, -5, 'bitcoin'],
        )).json()
        errors = ' '.join(bad['rows'][0]['errors'])

        for text in ('Ombor topilmadi', 'kelajakda', 'Mahsulot topilmadi', 'musbat',
                     'manfiy', 'To‘lov usuli'):
            self.assertIn(text, errors)


class SaleImportTests(ImportTestBase):

    def setUp(self):
        super().setUp()
        response = self.upload('purchases', 'commit', xlsx(
            ['Ombor', 'Sana', 'Artikul', 'Miqdor', 'Kirim narxi'],
            ['SHOP', self.today, 'CEM-1', 10, 50000],
        ))
        self.assertEqual(response.status_code, 201, response.content)

    def test_sales_use_card_price_and_check_stock(self):
        content = xlsx(
            ['Склад', 'Дата', 'Покупатель', 'Артикул', 'Количество', 'Цена продажи', 'Скидка, %'],
            ['SHOP', self.today, 'Chakana xaridor', 'CEM-1', 6, '', 10],
            ['SHOP', self.today, 'Chakana xaridor', 'CEM-1', 5, 60000, ''],
        )

        preview = self.upload('sales', 'preview', content).json()

        self.assertFalse(preview['can_commit'])
        self.assertEqual(preview['rows'][0]['errors'], [])
        # Ikkinchi qator birinchisidan qolgan 4 qopga sig'maydi
        self.assertIn('yetarli emas', preview['rows'][1]['errors'][0])

        ok = xlsx(
            ['Склад', 'Дата', 'Покупатель', 'Артикул', 'Количество', 'Цена продажи', 'Скидка, %'],
            ['SHOP', self.today, 'Chakana xaridor', 'CEM-1', 6, '', 10],
        )
        response = self.upload('sales', 'commit', ok)
        self.assertEqual(response.status_code, 201, response.content)

        with tenant_context(self.tenant.id):
            sale = Document.objects.get(pk=response.json()['documents'][0]['id'])
            self.assertEqual(sale.status, Document.Status.CONFIRMED)
            self.assertEqual(sale.customer_name, 'Chakana xaridor')
            self.assertEqual(sale.total_amount, D('313200.00'))
            total = StockBalance.objects.filter(variant=self.variant).aggregate(
                total=Sum('quantity')
            )['total']
            self.assertEqual(total, D('4'))

    def test_credit_sale_rejected(self):
        preview = self.upload('sales', 'preview', xlsx(
            ['Ombor', 'Sana', 'Artikul', 'Miqdor', 'To‘lov usuli'],
            ['SHOP', self.today, 'CEM-1', 1, 'Qarzga'],
        )).json()

        self.assertIn('Qarzga sotuvni', preview['rows'][0]['errors'][0])
