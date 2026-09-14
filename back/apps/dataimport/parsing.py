"""Excel faylni o'qish va katak qiymatlarini tahlil qilish.

Foydalanuvchi fayli hech qachon "toza" kelmaydi: sarlavha birinchi
qatorda bo'lmasligi, sonlar matn sifatida ("7 500 000,50"), sanalar
Excel seriya raqami sifatida (46279) yoki "14.09.2026" ko'rinishida
kelishi mumkin. Bu modul shu farqlarni bir joyda yutadi — import
qiluvchilar esa tayyor `Decimal` va `date` bilan ishlaydi.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

#: Yuklanadigan fayl chegarasi. 5000 qatorli to'la fayl ~1 MB bo'ladi.
MAX_FILE_SIZE = 5 * 1024 * 1024

#: Bir importdagi qatorlar chegarasi. Hammasi bitta tranzaksiyada
#: bajariladi — undan kattasi so'rovni daqiqalab ushlab turardi.
MAX_ROWS = 5000

#: Sarlavha qidiriladigan qatorlar soni: ba'zi fayllarda tepada
#: tashkilot nomi yoki sana yozilgan bo'ladi.
HEADER_SCAN_ROWS = 10

#: Excel sanalari shu kundan boshlab sanaladi (1900-yil kabisa xatosi
#: tufayli 31-dekabr emas, 30-dekabr).
EXCEL_EPOCH = date(1899, 12, 30)

#: O'zbekcha apostrof turli klaviaturada turlicha yoziladi
_UNIFY = str.maketrans({
    '‘': "'", '’': "'", 'ʻ': "'", 'ʼ': "'", '`': "'", '´': "'", 'ё': 'е',
})

_HEADER_NOISE = re.compile(r"[\s.,:;%()*№#_\-/'\"]+")


class ImportFileError(Exception):
    """Faylni umuman o'qib bo'lmaydi (format, hajm, bo'sh varaq)."""


def normalize_header(value) -> str:
    """Sarlavhani solishtirish kaliti: "Ед. изм." va "ед изм" bir xil."""
    text = str(value or '').lower().translate(_UNIFY)
    return _HEADER_NOISE.sub(' ', text).strip()


def name_key(value) -> str:
    """Nom bo'yicha qidirish kaliti: katta-kichik harf va bo'shliqsiz farq."""
    return ' '.join(str(value or '').translate(_UNIFY).split()).casefold()


def clean_text(value) -> str:
    """Katakni matnga aylantiradi.

    Excel uzun raqamlarni (shtrix-kod, SKU) son sifatida saqlaydi:
    4780000000011 `float` bo'lib keladi va `str()` uni "4780000000011.0"
    qilib qo'yardi.
    """
    if value is None:
        return ''

    if isinstance(value, bool):
        return 'ha' if value else "yo'q"

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    if isinstance(value, datetime):
        return value.strftime('%d.%m.%Y')

    if isinstance(value, date):
        return value.strftime('%d.%m.%Y')

    return ' '.join(str(value).split())


def parse_decimal(value) -> Decimal | None:
    """Sonni `Decimal` ga o'giradi. Bo'sh katak — `None`.

    Raises:
        ValueError: qiymat son emas.
    """
    if value is None or isinstance(value, bool):
        if isinstance(value, bool):
            raise ValueError(value)
        return None

    if isinstance(value, (int, Decimal)):
        return Decimal(value)

    if isinstance(value, float):
        # `repr` — eng qisqa aniq ko'rinish: 0.1 → "0.1", 0.1000000000000000055 emas
        return Decimal(repr(value))

    text = str(value).replace(' ', '').replace(' ', '').strip()

    for suffix in ("so'm", 'сўм', 'сум', 'uzs', '%'):
        text = re.sub(re.escape(suffix), '', text, flags=re.IGNORECASE)

    if not text:
        return None

    if ',' in text and '.' in text:
        # "7,500,000.50" — vergul ming ajratuvchi
        text = text.replace(',', '')
    elif text.count(',') > 1:
        text = text.replace(',', '')
    elif text.count('.') > 1:
        # "7.500.000" — nuqta ming ajratuvchi
        text = text.replace('.', '')
    else:
        # "7500000,50" — vergul o'nlik ajratuvchi
        text = text.replace(',', '.')

    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(value) from exc

    if not number.is_finite():
        raise ValueError(value)

    return number


_DATE_FORMATS = ('%d.%m.%Y', '%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d.%m.%y')


def parse_date(value) -> date | None:
    """Sanani o'qiydi. Bo'sh katak — `None`.

    Raises:
        ValueError: qiymat sana emas.
    """
    if value is None or value == '':
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Katak sana formatida bo'lmasa, Excel seriya raqamini beradi
        if 1 <= value < 100_000:
            return EXCEL_EPOCH + timedelta(days=int(value))
        raise ValueError(value)

    text = str(value).strip()

    if not text:
        return None

    # "2026-09-14 00:00:00" — vaqt qismini tashlaymiz
    text = text.split()[0]

    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    raise ValueError(value)


@dataclass
class SheetData:
    """O'qilgan varaq: sarlavhalar xaritasi va qatorlar."""

    name: str
    header_row: int
    #: Excel ustun tartibi → maydon kaliti
    columns: dict[int, str]
    #: Tanilmagan sarlavhalar (foydalanuvchiga ko'rsatish uchun)
    unknown: list[str] = field(default_factory=list)
    #: Tanilgan, lekin e'tiborsiz qoldirilgan sarlavhalar
    ignored: list[str] = field(default_factory=list)
    #: (Excel qator raqami, {kalit: xom qiymat})
    rows: list[tuple[int, dict]] = field(default_factory=list)
    file_hash: str = ''

    @property
    def keys(self) -> set[str]:
        return set(self.columns.values())


def read_sheet(upload, aliases: dict[str, str], ignored: frozenset[str] = frozenset()) -> SheetData:
    """Birinchi varaqni o'qiydi va sarlavhalarni maydonlarga bog'laydi.

    Arguments:
        upload: Django `UploadedFile`.
        aliases: normallashtirilgan sarlavha → maydon kaliti.
        ignored: tanilsa ham e'tiborsiz qoldiriladigan sarlavhalar
            (masalan boshqa tizim shablonidagi "Kompaniya").
    """
    if upload is None:
        raise ImportFileError('Fayl tanlanmagan.')

    if upload.size > MAX_FILE_SIZE:
        raise ImportFileError('Fayl hajmi 5 MB dan oshmasligi kerak.')

    content = upload.read()

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ImportFileError(
            'Faylni o‘qib bo‘lmadi. Excel (.xlsx) formatidagi fayl yuklang.'
        ) from exc

    try:
        if not workbook.worksheets:
            raise ImportFileError('Faylda varaq yo‘q.')

        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)

        head: list[tuple] = []

        for row in iterator:
            head.append(row)

            if len(head) >= HEADER_SCAN_ROWS:
                break

        header_index, columns, unknown, skipped = _find_header(head, aliases, ignored)

        if header_index is None:
            raise ImportFileError(
                'Sarlavha qatori topilmadi. Shablonni yuklab olib, ustun nomlarini '
                'o‘zgartirmasdan to‘ldiring.'
            )

        data = SheetData(
            name=sheet.title,
            header_row=header_index + 1,
            columns=columns,
            unknown=unknown,
            ignored=skipped,
            file_hash=hashlib.sha256(content).hexdigest(),
        )

        def collect(number: int, row: tuple) -> None:
            values = {
                key: row[index] if index < len(row) else None
                for index, key in columns.items()
            }

            if all(clean_text(value) == '' for value in values.values()):
                return

            if len(data.rows) >= MAX_ROWS:
                raise ImportFileError(
                    f'Bir faylda {MAX_ROWS} qatordan ko‘p bo‘lmasligi kerak. '
                    'Faylni bir necha qismga bo‘ling.'
                )

            data.rows.append((number, values))

        for offset, row in enumerate(head[header_index + 1:], start=header_index + 2):
            collect(offset, row)

        for number, row in enumerate(iterator, start=len(head) + 1):
            collect(number, row)

        return data
    finally:
        workbook.close()


def _find_header(head, aliases, ignored):
    """Eng ko'p tanilgan sarlavhali qatorni tanlaydi."""
    best = (None, {}, [], [])

    for index, row in enumerate(head):
        columns: dict[int, str] = {}
        unknown: list[str] = []
        skipped: list[str] = []

        for position, cell in enumerate(row):
            label = clean_text(cell)

            if not label:
                continue

            normalized = normalize_header(label)
            key = aliases.get(normalized)

            if key is None:
                (skipped if normalized in ignored else unknown).append(label)
                continue

            # Bir maydon ikki ustunda bo'lsa, birinchisi olinadi
            if key not in columns.values():
                columns[position] = key

        if len(columns) > len(best[1]):
            best = (index, columns, unknown, skipped)

    return best
