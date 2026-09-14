"""Excel import: shablon, oldindan ko'rish va saqlash.

    GET  /api/import/<tur>/template/  — bo'sh shablon va yo'riqnoma
    POST /api/import/<tur>/preview/   — fayl tekshiruvi, bazaga yozmaydi
    POST /api/import/<tur>/commit/    — tekshirib, hammasini bir tranzaksiyada saqlaydi

Tur: `categories`, `products`, `purchases`, `sales`. Har biri o'z bo'lim
ruxsatini talab qiladi (kategoriyalar, mahsulotlar, kirim, sotuv).
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass

from django.db import connection, transaction
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from rest_framework.exceptions import NotFound, PermissionDenied
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.audit.models import AuditEvent
from apps.core.export import HEADER_FILL, HEADER_FONT, excel_response
from apps.core.permissions import HasTenantMembership
from apps.dataimport.handlers import IMPORTERS, CategoryTree, CommitError, RowResult
from apps.dataimport.parsing import MAX_ROWS, ImportFileError, SheetData, normalize_header, read_sheet
from apps.dataimport.specs import SPECS, ImportSpec

TRUE_VALUES = {'1', 'true', 'yes', 'on'}


def _flag(value, default: bool) -> bool:
    if value in (None, ''):
        return default

    return str(value).strip().lower() in TRUE_VALUES


def find_duplicate(file_hash: str) -> dict | None:
    """Shu fayl avval import qilinganmi — tarixdagi yozuv bo'yicha."""
    event = (
        AuditEvent.objects.filter(action=AuditEvent.Action.IMPORT, changes__file_hash=file_hash)
        .order_by('-created_at')
        .first()
    )

    if event is None:
        return None

    return {
        'imported_at': event.created_at.isoformat(),
        'user_name': event.user_name,
        'object_repr': event.object_repr,
    }


@dataclass
class Prepared:
    spec: ImportSpec
    upload_name: str
    sheet: SheetData
    columns: list
    hidden_present: list[str]
    missing: list
    importer: object
    results: list[RowResult]
    duplicate: dict | None

    @property
    def valid(self) -> list[RowResult]:
        return [result for result in self.results if not result.errors]

    @property
    def has_errors(self) -> bool:
        return bool(self.missing) or len(self.valid) != len(self.results)

    @property
    def has_changes(self) -> bool:
        return any(result.action != 'skip' for result in self.valid)

    def payload(self) -> dict:
        valid = self.valid
        actions = Counter(result.action for result in valid)
        present = self.sheet.keys

        return {
            'type': self.spec.type,
            'title': self.spec.title,
            'file_name': self.upload_name,
            'sheet': self.sheet.name,
            'header_row': self.sheet.header_row,
            'columns': [
                {'key': column.key, 'label': column.label, 'required': column.required}
                for column in self.columns if column.key in present
            ],
            'missing_columns': [column.label for column in self.missing],
            'unknown_columns': self.sheet.unknown,
            'hidden_columns': self.hidden_present,
            'total': len(self.results),
            'valid': len(valid),
            'invalid': len(self.results) - len(valid),
            'with_warnings': sum(1 for result in self.results if result.warnings),
            'summary': {
                'create': actions['create'],
                'update': actions['update'],
                'skip': actions['skip'],
                'documents': self.importer.document_count(valid),
            },
            'duplicate': self.duplicate,
            'can_commit': not self.has_errors and self.has_changes,
            'rows': [result.as_dict() for result in self.results],
        }


class ImportBaseView(APIView):
    permission_classes = [HasTenantMembership]

    def get_spec(self, kind: str, *, write: bool) -> ImportSpec:
        spec = SPECS.get(kind)

        if spec is None:
            raise NotFound('Bunday import turi yo‘q.')

        membership = self.request.membership

        if not membership.has_perm(spec.permission):
            raise PermissionDenied('Bu bo‘limga ruxsatingiz yo‘q.')

        if write and not membership.can_write:
            raise PermissionDenied('Rolingiz ma’lumot kiritishga ruxsat bermaydi.')

        return spec

    def prepare(self, spec: ImportSpec, *, confirm: bool) -> Prepared:
        request = self.request
        upload = request.FILES.get('file')
        columns, hidden = spec.split_columns(request.membership)

        sheet = read_sheet(upload, spec.aliases(columns), spec.ignored_headers(hidden))

        hidden_aliases = spec.aliases(hidden) if hidden else {}
        hidden_present = sorted({
            hidden_aliases[normalize_header(label)] for label in sheet.ignored
            if normalize_header(label) in hidden_aliases
        })
        hidden_labels = [column.label for column in hidden if column.key in hidden_present]

        missing = [
            column for column in columns
            if column.required and column.key not in sheet.keys
        ]

        importer = IMPORTERS[spec.type](request, columns, confirm=confirm)
        results = [] if missing else importer.validate(sheet.rows)

        return Prepared(
            spec=spec,
            upload_name=getattr(upload, 'name', ''),
            sheet=sheet,
            columns=columns,
            hidden_present=hidden_labels,
            missing=missing,
            importer=importer,
            results=results,
            duplicate=find_duplicate(sheet.file_hash),
        )


class ImportTemplateView(ImportBaseView):

    def get(self, request, kind):
        spec = self.get_spec(kind, write=False)
        columns, _ = spec.split_columns(request.membership)

        return excel_response(self.build(spec, columns), f'{spec.filename}-shablon')

    def build(self, spec: ImportSpec, columns) -> Workbook:
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = spec.title[:31]

        for index, column in enumerate(columns, start=1):
            cell = sheet.cell(
                row=1, column=index, value=column.label + (' *' if column.required else '')
            )
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            sheet.column_dimensions[cell.column_letter].width = max(14, len(column.label) + 6)

        sheet.freeze_panes = 'A2'

        guide = workbook.create_sheet('Yo‘riqnoma')
        guide.append(['Ustun', 'Majburiy', 'Namuna', 'Izoh'])

        for cell in guide[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

        for column in columns:
            guide.append([
                column.label, 'ha' if column.required else '', column.sample, column.hint,
            ])

        guide.append([])

        for line in (
            'Birinchi varaq o‘qiladi. Sarlavhalarni o‘zgartirmang — ustunlar tartibi muhim emas.',
            f'Bir faylda {MAX_ROWS} qatorgacha. Bitta xatoli qator bo‘lsa ham fayl saqlanmaydi: '
            'avval tekshiruv natijasini ko‘rib, xatolarni tuzating.',
            'Ruscha (StoreFlow, 1C) sarlavhali fayllar ham qabul qilinadi.',
        ):
            guide.append([line])

        for row in self.reference_rows(spec):
            guide.append(row)

        for letter, width in zip('ABCD', (28, 10, 22, 70)):
            guide.column_dimensions[letter].width = width

        return workbook

    def reference_rows(self, spec: ImportSpec):
        """Faylni to'ldirishda kerak bo'ladigan mavjud qiymatlar."""
        if spec.type in ('purchases', 'sales'):
            from apps.warehouse.models import Warehouse, WarehouseAccess

            yield []
            yield ['Omborlar (kod — nomi)']

            for warehouse in WarehouseAccess.visible_to(
                self.request.user, Warehouse.objects.filter(is_active=True)
            ):
                yield [warehouse.code, '', '', warehouse.name]

        if spec.type == 'products':
            tree = CategoryTree()
            paths = sorted(
                ' > '.join(
                    node.name for node in _chain(tree, category)
                )
                for category in tree.by_id.values()
            )

            yield []
            yield ['Mavjud kategoriyalar']

            for path in paths[:500]:
                yield [path]


def _chain(tree: CategoryTree, category):
    nodes, node, seen = [], category, set()

    while node is not None and node.pk not in seen:
        seen.add(node.pk)
        nodes.append(node)
        node = tree.by_id.get(node.parent_id)

    return reversed(nodes)


class ImportPreviewView(ImportBaseView):
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, kind):
        spec = self.get_spec(kind, write=True)

        try:
            prepared = self.prepare(spec, confirm=_flag(request.data.get('confirm'), True))
        except ImportFileError as exc:
            return Response({'detail': str(exc)}, status=400)

        return Response(prepared.payload())


class ImportCommitView(ImportBaseView):
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, kind):
        spec = self.get_spec(kind, write=True)
        confirm = _flag(request.data.get('confirm'), True)
        allow_duplicate = _flag(request.data.get('allow_duplicate'), False)

        try:
            prepared = self.prepare(spec, confirm=confirm)
        except ImportFileError as exc:
            return Response({'detail': str(exc)}, status=400)

        if prepared.has_errors:
            return Response(
                {'detail': 'Faylda xatolar bor — tuzatib, qayta yuklang.', 'preview': prepared.payload()},
                status=400,
            )

        file_hash = prepared.sheet.file_hash

        try:
            with transaction.atomic():
                # Bir faylni ikki marta bosish yoki ikki oynadan bir vaqtda
                # yuborish: ikkinchisi birinchisi tugashini kutadi va keyin
                # uni tarixda ko'radi
                with connection.cursor() as cursor:
                    cursor.execute('SELECT pg_advisory_xact_lock(%s)', [int(file_hash[:15], 16)])

                duplicate = find_duplicate(file_hash)

                if duplicate and not allow_duplicate:
                    return Response(
                        {
                            'detail': 'Bu fayl avval import qilingan.',
                            'code': 'duplicate',
                            'duplicate': duplicate,
                        },
                        status=409,
                    )

                if not prepared.has_changes:
                    return Response(
                        {'detail': 'Faylda yangi yoki o‘zgargan ma’lumot yo‘q.'}, status=400
                    )

                outcome = prepared.importer.commit(
                    prepared.valid,
                    meta={
                        'filename': prepared.upload_name or 'fayl.xlsx',
                        'file_hash': file_hash,
                        'rows': len(prepared.valid),
                    },
                )
        except CommitError as exc:
            return Response({'detail': exc.messages}, status=400)

        return Response({'type': spec.type, 'confirmed': confirm, **outcome}, status=201)
