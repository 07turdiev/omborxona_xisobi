"""Hisobotni Excel faylga chiqarish.

Hisobot sahifasi bir necha bo'limdan iborat, shuning uchun fayl ham
**bir necha varaqdan** iborat bo'ladi: umumiy, kategoriya, ombor,
mahsulot, kunlik, yo'qotish, qoldiq qiymati. Ularni alohida fayllarga
ajratish buxgalterni har safar bir nechta faylni ochishga majbur
qilardi.

Xizmat funksiyalari `Decimal` qaytaradi va biz ularni serializerdan
o'tkazmaymiz — aks holda sonlar matnga aylanib, Excel'da qo'shilmay
qolardi.
"""

from __future__ import annotations

from typing import Any

from openpyxl import Workbook

from apps.core.export import (
    DATE,
    FORMATS,
    HEADER_FILL,
    HEADER_FONT,
    MONEY,
    NUMBER,
    QUANTITY,
    TEXT,
    Column,
    add_sheet,
)
from apps.core.access import Perm, visible_columns, visible_pairs
from apps.reports import services

#: Umumiy ko'rsatkichlar — "kalit, nom, format"
SUMMARY_LABELS = [
    ('purchase_count', 'Kirim hujjatlari', NUMBER),
    ('purchase_amount', 'Kirim summasi', MONEY),
    ('sale_count', 'Sotuv hujjatlari', NUMBER),
    ('revenue', 'Tushum', MONEY),
    ('cost', 'Tannarx (FIFO)', MONEY),
    ('gross_profit', 'Yalpi foyda', MONEY),
    ('margin_percent', 'Marja, %', MONEY),
    ('loss_amount', "Yo'qotishlar", MONEY),
    ('net_profit', 'Sof foyda', MONEY),
]

#: Qoldiq qiymati ko'rsatkichlari
VALUATION_LABELS = [
    ('positions', 'Pozitsiyalar', NUMBER),
    ('units', 'Jami miqdor', QUANTITY),
    ('reserved', 'Band qilingan', QUANTITY),
    ('cost_value', 'Tannarx qiymati', MONEY),
    ('retail_value', 'Sotuv qiymati', MONEY),
    ('potential_profit', 'Kutilayotgan foyda', MONEY),
    ('margin_percent', 'Marja, %', MONEY),
]

CATEGORY_COLUMNS = [
    Column('name', 'Kategoriya', TEXT, width=28),
    Column('quantity', 'Miqdor', QUANTITY),
    Column('revenue', 'Tushum', MONEY),
    Column('cost', 'Tannarx', MONEY),
    Column('profit', 'Foyda', MONEY),
    Column('margin_percent', 'Marja, %', MONEY),
]

WAREHOUSE_COLUMNS = [
    Column('name', 'Ombor', TEXT, width=28),
    Column('count', 'Hujjatlar', NUMBER),
    Column('revenue', 'Tushum', MONEY),
    Column('cost', 'Tannarx', MONEY),
    Column('profit', 'Foyda', MONEY),
]

PRODUCT_COLUMNS = [
    Column('name', 'Mahsulot', TEXT, width=32),
    Column('sku', 'SKU', TEXT, width=16),
    Column('quantity', 'Sotilgan miqdor', QUANTITY),
    Column('revenue', 'Tushum', MONEY),
    Column('cost', 'Tannarx', MONEY),
    Column('profit', 'Foyda', MONEY),
]

DAILY_COLUMNS = [
    Column('date', 'Sana', DATE),
    Column('count', 'Sotuvlar', NUMBER),
    Column('revenue', 'Tushum', MONEY),
    Column('profit', 'Foyda', MONEY),
]

LOSS_COLUMNS = [
    Column('label', 'Sababi', TEXT, width=28),
    Column('count', 'Hodisalar', NUMBER),
    Column('quantity', 'Miqdor', QUANTITY),
    Column('amount', 'Tannarx summasi', MONEY),
]


def _add_pairs_sheet(
    workbook: Workbook,
    labels: list[tuple[str, str, str]],
    data: dict,
    *,
    sheet_name: str,
    title: str,
    meta,
    sheet: Any = None,
):
    """«Ko'rsatkich / qiymat» ko'rinishidagi varaq.

    Bu yerda format ustun emas, **qator** bo'yicha aniqlanadi — har
    qatorda o'z turi bor. `add_sheet` ustun bo'yicha ishlagani uchun
    faqat sarlavhani undan olamiz, qatorlarni o'zimiz yozamiz.
    """
    sheet = add_sheet(
        workbook,
        [
            Column('label', "Ko'rsatkich", TEXT, width=28),
            Column('value', 'Qiymat', TEXT),
        ],
        [],
        sheet_name=sheet_name,
        title=title,
        meta=meta,
        sheet=sheet,
    )

    header_row = sheet.max_row

    for index, (key, label, kind) in enumerate(labels, start=header_row + 1):
        if key not in data:
            continue

        sheet.cell(row=index, column=1, value=label)

        value = data[key]
        cell = sheet.cell(
            row=index, column=2, value=float(value) if value is not None else None
        )
        cell.number_format = FORMATS[kind]

    # `add_sheet` bo'sh jadval uchun sarlavha uslubini qo'ygan, lekin
    # avtofiltrni qatorsiz belgilagan — qayta belgilaymiz
    for column in (1, 2):
        cell = sheet.cell(row=header_row, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    sheet.auto_filter.ref = f'A{header_row}:B{sheet.max_row}'

    return sheet


def build_report_workbook(
    period: dict, meta: list[tuple[str, str]], membership=None
) -> Workbook:
    """Butun hisobotni bitta ko'p varaqli kitobga yig'adi.

    `membership` berilsa, ruxsatsiz moliyaviy ustunlar faylga umuman
    chiqmaydi (bo'sh ustun emas — ustunning o'zi yo'q).
    """
    workbook = Workbook()

    def cols(columns):
        return visible_columns(columns, membership) if membership else columns

    def pairs(labels):
        return visible_pairs(labels, membership) if membership else labels

    show_loss_amount = membership is None or membership.has_perm(
        Perm.VIEW_PURCHASE_PRICE
    )

    losses = services.loss_summary(**period)

    _add_pairs_sheet(
        workbook,
        pairs(SUMMARY_LABELS),
        services.period_summary(**period),
        sheet_name='Umumiy',
        title='Umumiy ko‘rsatkichlar',
        meta=meta,
        sheet=workbook.active,
    )

    add_sheet(
        workbook,
        cols(CATEGORY_COLUMNS),
        services.by_category(**period),
        sheet_name='Kategoriya',
        title='Kategoriyalar bo‘yicha sotuv',
        meta=meta,
    )

    add_sheet(
        workbook,
        cols(WAREHOUSE_COLUMNS),
        services.by_warehouse(period['date_from'], period['date_to']),
        sheet_name='Omborlar',
        title='Omborlar bo‘yicha sotuv',
        meta=meta,
    )

    add_sheet(
        workbook,
        cols(PRODUCT_COLUMNS),
        # Eksportda ro'yxat cheklanmaydi: ekranda 10 ta yetadi, faylda
        # esa buxgalterga hammasi kerak bo'ladi.
        services.top_products(**period, limit=10000),
        sheet_name='Mahsulotlar',
        title='Mahsulotlar bo‘yicha sotuv',
        meta=meta,
    )

    add_sheet(
        workbook,
        cols(DAILY_COLUMNS),
        services.daily_sales(**period),
        sheet_name='Kunlik',
        title='Kunlik sotuv',
        meta=meta,
    )

    add_sheet(
        workbook,
        [c for c in LOSS_COLUMNS if show_loss_amount or c.key != 'amount'],
        losses['by_reason'],
        sheet_name='Yo‘qotishlar',
        title='Yo‘qotishlar sabablari bo‘yicha',
        meta=(
            [*meta, ('Jami', f"{losses['total']:,.2f}")] if show_loss_amount else meta
        ),
    )

    _add_pairs_sheet(
        workbook,
        pairs(VALUATION_LABELS),
        services.stock_valuation(period['warehouse']),
        sheet_name='Qoldiq qiymati',
        title='Qoldiq qiymati',
        meta=meta,
    )

    return workbook
